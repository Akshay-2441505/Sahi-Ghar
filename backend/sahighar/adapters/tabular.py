"""Read the tables an official data reply is likely to arrive in (CSV or Excel) into canonical columns.

Authorities word their column headers differently, so headers are matched by meaning (see _ALIASES),
not exactly. Values are parsed strictly: an unrecognised date is an error, never a guess.
"""
import csv
import io
import re
from calendar import month_abbr, month_name
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook


def _key(text: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text).lower()).strip()


# canonical column -> header wordings seen on MahaRERA's site or likely in a reply. Add wordings here.
_ALIASES = {
    "reg_no": ["project registration number", "project registration no", "registration number", "registration no",
               "rera registration number", "rera registration no", "rera no", "project no", "project number"],
    "name": ["project name", "name of project", "name"],
    "promoter_ref": ["promoter id", "promoter unique id", "promoter registration number", "promoter code", "promoter no"],
    "promoter_name": ["promoter name", "name of promoter", "promoter"],
    "city": ["district", "project district", "city"],
    "locality": ["locality", "village", "location"],
    "registration_end": ["registration valid up to", "registration valid till", "registration end date",
                         "original registration end date", "validity end date", "valid up to", "valid till",
                         "end date of original registration validity"],
    "extended_end": ["extended up to", "extension valid up to", "extension valid till", "extended end date",
                     "new valid up to", "revised end date", "extended till"],
    "address": ["registered office address", "registered address", "address", "registered office"],
    "pan": ["pan", "promoter pan", "pan number"],
    "partners": ["directors", "partners", "directors partners", "names of directors partners", "directors or partners"],
    "complaint_ref": ["complaint no", "complaint number", "complaint id"],
    "status": ["complaint status", "status"],
    "filed_date": ["date of filing", "filing date", "complaint date", "filed on", "date of complaint"],
    "filed_year": ["year of complaint filing", "filing year", "year"],
    "filed_month": ["month of complaint filing", "filing month", "month"],
    "non_execution": ["applied for non execution y n", "applied for non execution", "non execution applied",
                      "enforcement applied"],
    "order_url": ["order link", "order url"],
}
_CANON = {_key(alias): column for column, aliases in _ALIASES.items() for alias in aliases}


@dataclass
class Table:
    columns: list[str]  # canonical names, in file order
    rows: list[dict]  # canonical column -> str, or date for spreadsheet date cells
    lines: list[int]  # 1-based line/row number of each row in the original file, for error messages


def _csv_rows(data: bytes) -> list[list]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")  # common for spreadsheets exported on Windows in India
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text, newline=""), dialect))


def _xlsx_rows(data: bytes) -> list[list]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = [list(r) for r in workbook.worksheets[0].iter_rows(values_only=True)]  # first sheet only
    return [[c.date() if isinstance(c, datetime) else c for c in r] for r in rows]


def read_table(data: bytes, filename: str) -> Table:
    raw = _xlsx_rows(data) if filename.lower().endswith((".xlsx", ".xlsm")) else _csv_rows(data)
    start = next((i for i, r in enumerate(raw) if any(str(c or "").strip() for c in r)), None)
    if start is None:
        return Table([], [], [])
    mapping = {j: _CANON.get(_key(h)) for j, h in enumerate(raw[start])}
    if "complaint_ref" in mapping.values():  # in a complaints table "Project No." names the project
        mapping = {j: ("project_reg_no" if c == "reg_no" else c) for j, c in mapping.items()}
    rows, lines = [], []
    for offset, cells in enumerate(raw[start + 1:], start=start + 2):
        row = {}
        for j, column in mapping.items():
            if column is None or j >= len(cells):
                continue
            value = cells[j]
            row[column] = value if isinstance(value, date) else ("" if value is None else _text(value))
        if any(v != "" for v in row.values()):
            rows.append(row)
            lines.append(offset)
    return Table(list(dict.fromkeys(c for c in mapping.values() if c)), rows, lines)


def _text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # spreadsheet numbers arrive as floats: 2024.0 -> "2024"
    return str(value).strip()


def parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"unrecognised date {text!r} (expected DD/MM/YYYY)")


_MONTHS = {name.lower(): i for i, name in enumerate(month_name) if name} | {a.lower(): i for i, a in enumerate(month_abbr) if a}


def parse_month(value: object) -> int | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    if text in _MONTHS:
        return _MONTHS[text]
    if text.isdigit() and 1 <= int(text) <= 12:
        return int(text)
    raise ValueError(f"unrecognised month {text!r}")


def parse_year(value: object) -> int | None:
    text = re.sub(r"\.0+$", "", _text(value)) if value not in (None, "") else ""  # "2024.0" from spreadsheets
    if not text:
        return None
    if text.isdigit() and 1990 <= int(text) <= 2100:
        return int(text)
    raise ValueError(f"unrecognised year {text!r} (expected four digits)")


def parse_yes_no(value: object) -> bool:
    text = str(value or "").strip().lower()
    if text in ("y", "yes", "true", "1"):
        return True
    if text in ("n", "no", "false", "0", ""):
        return False
    raise ValueError(f"expected Y or N, got {text!r}")


def split_names(value: object) -> list[str]:
    """Directors or partners in one cell: split on ; | or new lines, never on commas ("Shah, Ramesh")."""
    return [n.strip() for n in re.split(r"[;|\n]", str(value or "")) if n.strip()]
